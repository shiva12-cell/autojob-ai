import pytest
import time
import openpyxl
from starlette.testclient import TestClient
from src.web.server import app
from src.excel_manager import ExcelManager
from src.database.db import db

def test_excel_ensure_custom_role_sheet(tmp_path):
    excel_path = tmp_path / "test_tracker.xlsx"
    mgr = ExcelManager(file_path=str(excel_path))
    
    # 1. Ensure a custom role sheet is created only when requested
    sheet_name = mgr.ensure_custom_role_sheet("React Developer")
    assert sheet_name == "React Developer"
    
    wb = openpyxl.load_workbook(str(excel_path))
    assert "React Developer" in wb.sheetnames
    ws = wb["React Developer"]
    # Check headers
    headers = [cell.value for cell in ws[1]]
    assert "Application ID" in headers
    assert "Job Title" in headers
    assert "Company" in headers
    wb.close()
    
    # 2. Test sanitize and max 31 characters
    long_role = "Senior Full Stack Cloud Solutions Architect & Engineer"
    sanitized_name = mgr.ensure_custom_role_sheet(long_role)
    assert len(sanitized_name) <= 31
    wb = openpyxl.load_workbook(str(excel_path))
    assert sanitized_name in wb.sheetnames
    wb.close()

def test_api_keywords_crud_and_sheet_creation(tmp_path):
    client = TestClient(app)
    
    # 1. Fetch default keywords
    res = client.get("/api/keywords")
    assert res.status_code == 200
    roles = res.json()
    role_labels = [r["label"].lower() for r in roles]
    assert "data analyst" in role_labels
    assert "business analyst" in role_labels
    
    # 2. Add custom role (multi-word)
    custom_role = "Senior React Lead"
    res_add = client.post("/api/keywords", json={"role": custom_role})
    assert res_add.status_code == 200
    data_add = res_add.json()
    assert data_add["status"] == "success"
    assert data_add["role"] == custom_role
    assert "sheet_name" in data_add
    
    # 3. Verify it shows up in GET /api/keywords with is_custom=True
    res_get = client.get("/api/keywords")
    assert res_get.status_code == 200
    updated_roles = res_get.json()
    matched = [r for r in updated_roles if r["label"] == custom_role]
    assert len(matched) == 1
    assert matched[0]["is_custom"] is True
    
    # 4. Update the role
    new_role = "Staff React Lead"
    res_put = client.put("/api/keywords", json={"old_role": custom_role, "new_role": new_role})
    assert res_put.status_code == 200
    data_put = res_put.json()
    assert data_put["status"] == "success"
    
    # 5. Delete the role using query parameter or json
    res_del = client.delete(f"/api/keywords?role={new_role}")
    assert res_del.status_code == 200
    data_del = res_del.json()
    assert data_del["status"] == "success"

def test_database_target_roles_crud():
    user_id = f"test_role_user_{int(time.time()*1000)}"
    user = db.create_user(user_id, f"{user_id}@domain.com", "hash1", "salt1")
    
    # Initial target roles
    initial_roles = db.get_target_roles(user_id)
    assert "Data Analyst" in initial_roles
    
    # Add new role (one or many words)
    roles = db.add_target_role(user_id, "Principal AI Scientist")
    assert "Principal AI Scientist" in roles
    
    # Update role
    roles = db.update_target_role(user_id, "Principal AI Scientist", "Chief AI Scientist")
    assert "Chief AI Scientist" in roles
    assert "Principal AI Scientist" not in roles
    
    # Delete role
    roles = db.delete_target_role(user_id, "Chief AI Scientist")
    assert "Chief AI Scientist" not in roles
